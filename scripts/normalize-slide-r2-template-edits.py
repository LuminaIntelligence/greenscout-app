#!/usr/bin/env python3
"""One-shot helper: consolidated template edits for Runde-2 defects.

Defekte R2-4, R2-6, R2-8, R2-10 (2026-05-31) — see DECISIONS.md.

This script applies four orthogonal template-edit fixes in a single
pass so the binary commit stays atomic. Each section is idempotent.

R2-4 — Slide 4 headline numbers TEXT_TO_FIT_SHAPE
    The headline-number shapes on slide 4 (Anlage kWp, Jahresertrag
    kWh, Eigenverbrauchsquote %, Stromersparnis 20-Jahre €,
    Pacht-Einmalzahlung €) have ``auto_size = None`` (i.e. inherit
    the layout default SHAPE_TO_FIT_TEXT). When the value is large
    enough — e.g. ``3.000.000`` for ``pv_erzeugung_kwh_jahr`` — the
    text grows past the shape's bounds and collides with the static
    "kWh" / "€" label sitting next to it. TEXT_TO_FIT_SHAPE shrinks
    the font instead. Five shapes affected — see SLIDE_4_HEADLINE_*.

R2-6 — Slide 16 Variantenspalten TEXT_TO_FIT_SHAPE (verification)
    PR #54 (Defekt C4) already set TEXT_TO_FIT_SHAPE on the
    Variantenvergleich shapes. After PR #55 (D1+D2+D3) introduced
    longer phrase-text in some bullets, defense-in-depth re-applies
    the same shapes here. Idempotent.

R2-8 — Slide 1 Berater-Name TEXT_TO_FIT_SHAPE (verification)
    PR #54 (Defekt C3) already set TEXT_TO_FIT_SHAPE on the slide-1
    Berater shape. We re-verify here in case a later commit reverted
    it; if the property is already set this section is a no-op.

R2-10 — Slide 17 column heights uniform (verification)
    PR #53 (Defekt C2) normalised the slide-17 grid to identical
    Y-positions and a uniform Inhalte-height. We re-verify those
    invariants here; if any height drift is detected we re-apply
    the normalised height.

Run once from the repo root and commit the resulting template binary.

Usage::

    python scripts/normalize-slide-r2-template-edits.py [path/to/template.pptx]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

DEFAULT_TEMPLATE = Path("templates/Machbarkeitsstudie-PV-Template_v1_6.pptx")

# --- R2-4 — Slide 4 headline-number shapes -----------------------------
#
# Identified via scripts/inspect-pptx.py on 2026-05-31. The five big
# numeric shapes that risk overflowing their bounding box when the
# dynamic value runs long. Each contains ONLY a placeholder token (or
# a token + tiny suffix like "%") — i.e. these are not body-text
# shapes where TEXT_TO_FIT_SHAPE would shrink legibility-critical
# explanatory text.
SLIDE_4_INDEX = 3
SLIDE_4_HEADLINE_SHAPE_IDS = {
    6,   # "Text 4"  — {{anlage_kwp}} (kWp value)
    9,   # "Text 7"  — {{pv_erzeugung_kwh_jahr}} (Jahresertrag — the "3.000.000" overflow)
    15,  # "Text 13" — {{eigenverbrauchsquote_prozent}}% (Eigenverbrauch %)
    19,  # "Text 16" — {{ersparnis_gesamt_vertragslaufzeit_eur}} (20-Jahres-Ersparnis €)
    29,  # "Text 16" — {{pacht_einnahme_einmalig_eur}} (Pacht-Einmalzahlung €)
}

# --- R2-6 — Slide 16 Varianten shapes (re-application, idempotent) ----
SLIDE_16_INDEX = 15
SLIDE_16_VARIANTE_SHAPE_NAMES = {
    "Text 2", "Text 3", "Text 4", "Text 5", "Text 6", "Text 7",
    "Text 8", "Text 9", "Text 10", "Text 11", "Text 12", "Text 13",
}

# --- R2-8 — Slide 1 Berater-Name shape (re-verification, idempotent) --
SLIDE_1_INDEX = 0
SLIDE_1_BERATER_SHAPE_NAME = "Textfeld 3"

# --- R2-10 — Slide 17 column-height + top-position verification -------
#
# PR #53 normalised the slide-17 grid. R2-10 (Runde 2) re-verifies
# both rows' heights AND tops for drift, since the live customer
# template showed the Ergebnisse-row column-5 "sitting lower" than
# the rest. Inventory captured 2026-05-31 — 7 columns, not 6 as the
# Runde-2 report assumed (Textfeld 20 is the 4th column,
# left-to-right by L-position).
SLIDE_17_INDEX = 16

#: Inhalte shapes (upper row): re-verify uniform height + top.
SLIDE_17_INHALTE_NAMES = {
    "Textfeld 2", "Textfeld 3", "Textfeld 4", "Textfeld 5",
    "Textfeld 6", "Textfeld 7", "Textfeld 20",
}

#: Ergebnisse shapes (lower row): re-verify uniform top across all 7
#: columns. Heights are intentionally variable (each box contains
#: differently-sized text); the column-5-sitting-lower defect is
#: about the TOP alignment, not the height.
SLIDE_17_ERGEBNISSE_NAMES = {
    "Textfeld 8", "Textfeld 9", "Textfeld 10", "Textfeld 11",
    "Textfeld 12", "Textfeld 13", "Textfeld 21",
}


def _apply_fit_to_shape(shape: Any, changes: list[str]) -> None:
    """Apply TEXT_TO_FIT_SHAPE + word_wrap=True to ``shape``, recording
    each actual mutation in ``changes`` (no-op-safe if already set).
    """
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if tf.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
        old = tf.auto_size
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        changes.append(f"auto_size {old!r}→TEXT_TO_FIT_SHAPE")
    if not tf.word_wrap:
        tf.word_wrap = True
        changes.append("word_wrap→True")


def _normalize_slide_4(slide: Any, modified: list[tuple[str, str, int, str]]) -> int:
    """R2-4: TEXT_TO_FIT_SHAPE on slide-4 headline numbers."""
    seen: set[int] = set()
    for shape in slide.shapes:
        sid = int(shape.shape_id)
        if sid not in SLIDE_4_HEADLINE_SHAPE_IDS:
            continue
        seen.add(sid)
        changes: list[str] = []
        _apply_fit_to_shape(shape, changes)
        modified.append(
            (
                "R2-4 slide4-headline",
                shape.name,
                sid,
                ", ".join(changes) or "noop",
            )
        )
    missing = SLIDE_4_HEADLINE_SHAPE_IDS - seen
    if missing:
        print(
            f"ERROR: slide-4 headline shape ids not found: {sorted(missing)}. "
            "Re-run inspect-pptx and update SLIDE_4_HEADLINE_SHAPE_IDS.",
            file=sys.stderr,
        )
        return -1
    return 0


def _normalize_slide_16(slide: Any, modified: list[tuple[str, str, int, str]]) -> int:
    """R2-6: re-apply TEXT_TO_FIT_SHAPE on slide-16 Variantenspalten."""
    seen: set[str] = set()
    for shape in slide.shapes:
        if shape.name not in SLIDE_16_VARIANTE_SHAPE_NAMES:
            continue
        seen.add(shape.name)
        changes: list[str] = []
        _apply_fit_to_shape(shape, changes)
        modified.append(
            (
                "R2-6 slide16-variante",
                shape.name,
                int(shape.shape_id),
                ", ".join(changes) or "noop",
            )
        )
    missing = SLIDE_16_VARIANTE_SHAPE_NAMES - seen
    if missing:
        print(
            f"ERROR: slide-16 variante shapes not found: {sorted(missing)}.",
            file=sys.stderr,
        )
        return -1
    return 0


def _normalize_slide_1(slide: Any, modified: list[tuple[str, str, int, str]]) -> int:
    """R2-8: re-apply TEXT_TO_FIT_SHAPE on slide-1 Berater-Name shape."""
    for shape in slide.shapes:
        if shape.name != SLIDE_1_BERATER_SHAPE_NAME:
            continue
        changes: list[str] = []
        _apply_fit_to_shape(shape, changes)
        modified.append(
            (
                "R2-8 slide1-berater",
                shape.name,
                int(shape.shape_id),
                ", ".join(changes) or "noop",
            )
        )
        return 0
    print(
        f"ERROR: slide-1 shape {SLIDE_1_BERATER_SHAPE_NAME!r} not found.",
        file=sys.stderr,
    )
    return -1


def _snap_uniform(
    shapes: list[Any],
    attr: str,
    modified: list[tuple[str, str, int, str]],
    label: str,
) -> None:
    """Snap an attribute (e.g. 'top', 'height') across shapes to a uniform value.

    Picks the most-common existing value as canonical and applies it
    to every outlier. Idempotent — no-op when all values already match.
    """
    from collections import Counter

    values = [int(getattr(s, attr)) for s in shapes if getattr(s, attr) is not None]
    if not values:
        return
    canonical = Counter(values).most_common(1)[0][0]
    for shape in shapes:
        current = int(getattr(shape, attr)) if getattr(shape, attr) is not None else 0
        if current != canonical:
            setattr(shape, attr, canonical)
            modified.append(
                (
                    label,
                    shape.name,
                    int(shape.shape_id),
                    f"{attr} {current}→{canonical}",
                )
            )
        else:
            modified.append(
                (
                    label,
                    shape.name,
                    int(shape.shape_id),
                    f"{attr} noop (already {canonical})",
                )
            )


def _normalize_slide_17(slide: Any, modified: list[tuple[str, str, int, str]]) -> int:
    """R2-10: verify slide-17 grid Inhalte heights + Ergebnisse tops.

    PR #53 normalised the grid the first time. R2-10 asserts the
    invariants are still in place AND fixes any drift detected on:

    - Inhalte row: uniform ``height`` (all 7 boxes same size).
    - Inhalte row: uniform ``top`` (all 7 sit on the same baseline).
    - Ergebnisse row: uniform ``top`` (the "Spalte 5 sitzt tiefer"
      defect that triggered Runde 2 — heights stay variable to fit
      the per-column wording).
    """
    inhalte_shapes = [s for s in slide.shapes if s.name in SLIDE_17_INHALTE_NAMES]
    if inhalte_shapes:
        _snap_uniform(inhalte_shapes, "height", modified, "R2-10 slide17-inhalte")
        _snap_uniform(inhalte_shapes, "top", modified, "R2-10 slide17-inhalte")
    else:
        print(
            "WARN: no slide-17 Inhalte shapes found — skipping height/top verification.",
            file=sys.stderr,
        )

    ergebnisse_shapes = [s for s in slide.shapes if s.name in SLIDE_17_ERGEBNISSE_NAMES]
    if ergebnisse_shapes:
        _snap_uniform(ergebnisse_shapes, "top", modified, "R2-10 slide17-ergebnisse")
    else:
        print(
            "WARN: no slide-17 Ergebnisse shapes found — skipping top verification.",
            file=sys.stderr,
        )
    return 0


def normalize(template_path: Path) -> int:
    """Apply all four template-edit fixes. Returns 0 on success, <0 on error."""
    pres = Presentation(str(template_path))
    if len(pres.slides) <= SLIDE_17_INDEX:
        print(f"ERROR: template has fewer than 17 slides", file=sys.stderr)
        return -1

    modified: list[tuple[str, str, int, str]] = []

    rc = _normalize_slide_4(pres.slides[SLIDE_4_INDEX], modified)
    if rc < 0:
        return rc
    rc = _normalize_slide_16(pres.slides[SLIDE_16_INDEX], modified)
    if rc < 0:
        return rc
    rc = _normalize_slide_1(pres.slides[SLIDE_1_INDEX], modified)
    if rc < 0:
        return rc
    rc = _normalize_slide_17(pres.slides[SLIDE_17_INDEX], modified)
    if rc < 0:
        return rc

    if not modified:
        print("WARN: nothing to normalize.")
        return 0

    # Save only if there was at least one real change (not just no-ops).
    real_changes = [m for m in modified if "noop" not in m[3]]
    if real_changes:
        pres.save(str(template_path))
        print(f"Modified {len(real_changes)} shape(s) (of {len(modified)} visited):")
        for kind, name, sid, change in modified:
            marker = "  " if "noop" in change else " *"
            print(f"  {marker} {kind:24s} {name:14s} (id={sid:3d})  {change}")
    else:
        print(f"All {len(modified)} shape(s) already normalised — template unchanged.")
        for kind, name, sid, change in modified:
            print(f"    {kind:24s} {name:14s} (id={sid:3d})  {change}")
    return 0


def main() -> int:
    """Entrypoint."""
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
        sys.stderr.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, OSError):
        pass

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "template",
        type=Path,
        nargs="?",
        default=DEFAULT_TEMPLATE,
        help="Path to the .pptx template (defaults to the v1.6 GreenScout template).",
    )
    args = parser.parse_args()

    if not args.template.exists():
        print(f"ERROR: template not found at {args.template}", file=sys.stderr)
        return 2

    rc = normalize(args.template)
    if rc < 0:
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main())
