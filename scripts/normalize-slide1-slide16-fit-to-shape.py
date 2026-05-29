#!/usr/bin/env python3
"""One-shot helper: TEXT_TO_FIT_SHAPE on Slide 1 Berater-Shape and
Slide 16 Variantenspalten.

Defekte C3 + C4 (2026-05-29) — see DECISIONS.md.

Slide 1 (Defekt C3): ``Textfeld 3`` renders the line
"Eingereicht über {{consultant_full_name}} / direkt vom Unternehmen".
The original template ships with ``auto_size = SHAPE_TO_FIT_TEXT``, so
the shape grows vertically with the text. With a long Berater-Name
(e.g. "Admin GreenScout" — 16 chars vs. the template-example
"Bernd Berater" — 13 chars) the single-line shape would have to grow
downward past the visible slide region, and the trailing word
"Unternehmen" gets clipped. Switching to ``TEXT_TO_FIT_SHAPE`` shrinks
the font instead, keeping the full sentence visible.

Slide 16 (Defekt C4): Variantenvergleich grid (two columns × 6 rows of
bullets/titles). Every shape ships with ``auto_size = None`` — text
silently overflows the shape bounds. After PR #50 (Pacht-Formel A1 fix)
the Pacht value shrank from ~1.000.000 € to ~50.000 €, so the immediate
mid-word truncation ("…ca.1" instead of "…ca. 156.000 €") is unlikely
to recur today, but defense-in-depth ``TEXT_TO_FIT_SHAPE`` on every
Varianten-shape future-proofs the layout against longer dynamic values
(higher Ersparnis, longer object names, etc.).

Affected shapes:

- Slide 1: ``Textfeld 3``
- Slide 16 Variante A column: ``Text 2`` (Titel), ``Text 3``–``Text 7``
- Slide 16 Variante B column: ``Text 8`` (Titel), ``Text 9``–``Text 13``

The slide-16 title (``Text 0``), the customer-object subtitle
(``Text 1``) and the page-number placeholder are intentionally left
alone — they are static layout chrome, not Varianten-data shapes.

Run once from the repo root and commit the resulting template binary.

Usage::

    python scripts/normalize-slide1-slide16-fit-to-shape.py [path/to/template.pptx]

Default path is ``templates/Machbarkeitsstudie-PV-Template_v1_6.pptx``
relative to the repo root.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

DEFAULT_TEMPLATE = Path("templates/Machbarkeitsstudie-PV-Template_v1_6.pptx")

SLIDE_1_INDEX = 0
SLIDE_16_INDEX = 15

# Shape names identified via scripts/inspect-pptx.py on 2026-05-29
# (see DECISIONS C3/C4 entry for the full inventory).
SLIDE_1_BERATER_SHAPE_NAME = "Textfeld 3"

# Slide-16 Varianten columns (Titel + bullet body shapes).
# Variante A: Text 2 (Titel), Text 3..7 (bullets)
# Variante B: Text 8 (Titel), Text 9..13 (bullets)
# Text 0 (slide title) and Text 1 (object subtitle) are layout chrome
# and intentionally excluded.
SLIDE_16_VARIANTE_SHAPE_NAMES = {
    "Text 2",
    "Text 3",
    "Text 4",
    "Text 5",
    "Text 6",
    "Text 7",
    "Text 8",
    "Text 9",
    "Text 10",
    "Text 11",
    "Text 12",
    "Text 13",
}


def _apply_fit_to_shape(shape, changes: list[str]) -> None:
    """Apply TEXT_TO_FIT_SHAPE + word_wrap=True to ``shape``, recording
    each actual mutation in ``changes`` (no-op-safe if already set).
    """
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if tf.auto_size != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE:
        old = tf.auto_size
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        changes.append(f"auto_size {old!r}→TEXT_TO_FIT_SHAPE")
    if not tf.word_wrap:
        tf.word_wrap = True
        changes.append("word_wrap→True")


def normalize(template_path: Path) -> int:
    """Apply the two-slide fix. Returns the count of matched shapes (>=0)
    or a negative error code.
    """
    pres = Presentation(str(template_path))
    if len(pres.slides) <= SLIDE_16_INDEX:
        print(
            f"ERROR: slide {SLIDE_16_INDEX + 1} does not exist in {template_path}",
            file=sys.stderr,
        )
        return -1

    modified: list[tuple[str, str, int, str]] = []  # (kind, name, sid, change)

    # --- Slide 1 ------------------------------------------------------
    slide1 = pres.slides[SLIDE_1_INDEX]
    seen_slide1 = False
    for shape in slide1.shapes:
        if shape.name == SLIDE_1_BERATER_SHAPE_NAME:
            seen_slide1 = True
            changes: list[str] = []
            _apply_fit_to_shape(shape, changes)
            modified.append(
                (
                    "slide1-berater",
                    shape.name,
                    int(shape.shape_id),
                    ", ".join(changes) or "noop",
                )
            )
            break
    if not seen_slide1:
        print(
            f"ERROR: Slide-1 shape {SLIDE_1_BERATER_SHAPE_NAME!r} not found — "
            "re-run inspect-pptx and update the constant.",
            file=sys.stderr,
        )
        return -2

    # --- Slide 16 -----------------------------------------------------
    slide16 = pres.slides[SLIDE_16_INDEX]
    seen_slide16: set[str] = set()
    for shape in slide16.shapes:
        if shape.name in SLIDE_16_VARIANTE_SHAPE_NAMES:
            seen_slide16.add(shape.name)
            changes = []
            _apply_fit_to_shape(shape, changes)
            modified.append(
                (
                    "slide16-variante",
                    shape.name,
                    int(shape.shape_id),
                    ", ".join(changes) or "noop",
                )
            )

    missing = SLIDE_16_VARIANTE_SHAPE_NAMES - seen_slide16
    if missing:
        print(
            "ERROR: expected Varianten-shapes not found on slide 16: "
            f"{sorted(missing)}. Update the constants and re-run inspect-pptx.",
            file=sys.stderr,
        )
        return -3

    if not modified:
        print("WARN: no shapes matched — nothing to normalize.")
        return 0

    pres.save(str(template_path))
    print(f"Modified {len(modified)} shapes:")
    for kind, name, sid, change in modified:
        print(f"  {kind:18s} {name:14s} (id={sid:3d})  {change}")
    return len(modified)


def main() -> int:
    """Entrypoint."""
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
        sys.stderr.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, OSError):
        pass

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "template",
        type=Path,
        nargs="?",
        default=DEFAULT_TEMPLATE,
        help="Path to the .pptx template (defaults to the v1.6 GreenScout template).",
    )
    args = parser.parse_args()

    if not args.template.exists():
        print(f"ERROR: template not found at {args.template}", file=sys.stderr)
        return 2

    result = normalize(args.template)
    if result < 0:
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main())
