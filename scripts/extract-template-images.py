"""
Pivot-2b PASS 3 — Mechanische Bild-Extraktion aus dem PPTX-Original.

Iteriert alle Slides, identifiziert Picture-Shapes (inkl. verschachtelte
Group-Shapes) und schreibt das eingebettete Bild-Blob nach
`public/assets/pptx-slide{N}-image{i}.{ext}`. Output: `public/assets/`,
plus eine Manifest-Datei `public/assets/pptx-images-manifest.json` mit
Slide-Nr, Shape-Name, Größe (EMU), Dateiname und Bytes.

Pendant zu `scripts/extract-template-text.py`. Nutzt das selbe
mini-venv `scripts/.venv-pptx-extract/` (siehe
`scripts/README-extract-template.md`). Sicher gegen Drift:

- Bilder werden NICHT manuell zugeordnet — sie kommen aus diesem
  Skript. Wenn das PPTX aktualisiert wird, läuft dieses Skript erneut
  und der Diff im Review zeigt was sich bildlich geändert hat.
- Das eingebettete Blob ist die Source of Truth; die Reproduktion soll
  KEIN externes Asset einbinden, das nicht aus dem PPTX kommt.

Wann erneut laufen? Nach jedem PPTX-Template-Update.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

REPO_ROOT = Path(__file__).resolve().parent.parent
PPTX_PATH = REPO_ROOT / "templates" / "Machbarkeitsstudie-PV-Template_v1_6.pptx"
OUT_DIR = REPO_ROOT / "public" / "assets"
MANIFEST_PATH = OUT_DIR / "pptx-images-manifest.json"


def iter_picture_shapes(shapes, slide_number: int, path_prefix: str = ""):
    """Yields (slide_number, shape, full_path_string) for every PICTURE-Shape.

    Recursively descends into GROUP-Shapes. The path string is a `/`-joined
    breadcrumb of shape names so deeply nested pictures get a unique
    identifier in the manifest.
    """
    for shape in shapes:
        breadcrumb = f"{path_prefix}/{shape.name}" if path_prefix else shape.name
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_picture_shapes(shape.shapes, slide_number, breadcrumb)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield slide_number, shape, breadcrumb


def extract() -> None:
    if not PPTX_PATH.exists():
        raise SystemExit(
            f"PPTX not found at {PPTX_PATH}. Restore it from git history via "
            "`git show <pre-pivot-sha>:templates/Machbarkeitsstudie-PV-Template_v1_6.pptx > "
            "templates/Machbarkeitsstudie-PV-Template_v1_6.pptx` (see "
            "scripts/README-extract-template.md)."
        )

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(PPTX_PATH))
    manifest: list[dict] = []
    skipped: list[dict] = []
    written = 0
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for picture_idx, (sn, shape, breadcrumb) in enumerate(
            iter_picture_shapes(slide.shapes, slide_idx), start=1
        ):
            # Picture shapes can be EMBEDDED (blob in /ppt/media) or LINKED
            # (external path reference). Linked pictures don't carry a blob;
            # record them in `skipped` for traceability and move on.
            try:
                image = shape.image
            except ValueError as exc:
                skipped.append(
                    {
                        "slide_number": sn,
                        "picture_index": picture_idx,
                        "shape_name": shape.name,
                        "shape_id": shape.shape_id,
                        "shape_path": breadcrumb,
                        "reason": str(exc),
                    }
                )
                continue
            blob: bytes = image.blob
            ext = image.ext.lower()
            filename = f"pptx-slide{sn:02d}-image{picture_idx}.{ext}"
            target = OUT_DIR / filename
            target.write_bytes(blob)
            written += 1
            manifest.append(
                {
                    "slide_number": sn,
                    "picture_index": picture_idx,
                    "shape_name": shape.name,
                    "shape_id": shape.shape_id,
                    "shape_path": breadcrumb,
                    "filename": filename,
                    "ext": ext,
                    "byte_size": len(blob),
                    "sha256_hex": hashlib.sha256(blob).hexdigest(),
                    "left_emu": shape.left,
                    "top_emu": shape.top,
                    "width_emu": shape.width,
                    "height_emu": shape.height,
                }
            )

    MANIFEST_PATH.write_text(
        json.dumps(
            {"images": manifest, "skipped_linked_pictures": skipped},
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    by_slide: dict[int, int] = {}
    for entry in manifest:
        by_slide[entry["slide_number"]] = by_slide.get(entry["slide_number"], 0) + 1
    print(
        f"Wrote {written} images to {OUT_DIR.relative_to(REPO_ROOT)} "
        f"(manifest: {MANIFEST_PATH.relative_to(REPO_ROOT)})."
    )
    for sn in sorted(by_slide):
        print(f"  slide {sn:02d}: {by_slide[sn]} embedded image(s)")
    if skipped:
        print(f"Skipped {len(skipped)} linked-only picture(s) — see manifest.skipped_linked_pictures.")


if __name__ == "__main__":
    extract()
