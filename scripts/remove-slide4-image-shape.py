#!/usr/bin/env python3
"""One-shot helper: remove the rogue `image_before` shape on slide 4.

Defekt B1 (2026-05-29) — see DECISIONS.md.

Background: Slice-3b's T-037 placeholder migration renamed `Image 0` on
slide 4 to `image_before` based on disambiguation Q5 with the caveat
"visuell prüfen". The first production-generated PPTX revealed the
mistake: slide 4's "Eigenverbrauch" composition (Shapes 13/14/15/17
etc.) has no photo slot in the original template — the renamed shape
sat directly over the `4 %` headline, obscuring the graphic and
leaving the "Eigenverbrauch"-label without anchor.

Fix: drop the shape from slide 4 entirely. Slide 5's
`image_before` / `image_after` mapping (`Grafik 2` / `Grafik 5`) is
unaffected and remains the only place BEFORE / AFTER photos appear.

Run once from the repo root. The template is the source of truth after
this fix lands — `apply-pptx-placeholders.py` no longer recreates the
shape (its `IMAGE_RENAMES` entry was removed in the same PR).

Usage::

    python scripts/remove-slide4-image-shape.py [path/to/template.pptx]

Default path is ``templates/Machbarkeitsstudie-PV-Template_v1_6.pptx``
relative to the repo root.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation

DEFAULT_TEMPLATE = Path("templates/Machbarkeitsstudie-PV-Template_v1_6.pptx")
SLIDE_INDEX = 3  # slide 4 (0-indexed)
SHAPE_NAME = "image_before"


def remove_slide4_image(template_path: Path) -> int:
    """Remove every `image_before` shape from slide 4. Returns the count removed."""
    pres = Presentation(str(template_path))
    if SLIDE_INDEX >= len(pres.slides):
        print(
            f"ERROR: slide {SLIDE_INDEX + 1} does not exist in {template_path}",
            file=sys.stderr,
        )
        return -1

    slide4 = pres.slides[SLIDE_INDEX]
    removed: list[int] = []
    # Iterate over a snapshot because we mutate the underlying tree.
    for shape in list(slide4.shapes):
        if shape.name == SHAPE_NAME:
            sp = shape._element  # type: ignore[attr-defined]  # python-pptx exposes no public delete API.
            sp.getparent().remove(sp)
            removed.append(int(shape.shape_id))
            print(
                f"Removed shape '{SHAPE_NAME}' (id={shape.shape_id}) from slide {SLIDE_INDEX + 1}"
            )

    if not removed:
        print(
            f"WARN: no shape named '{SHAPE_NAME}' found on slide {SLIDE_INDEX + 1} — "
            "template may already be clean.",
            file=sys.stderr,
        )
        return 0

    pres.save(str(template_path))
    print(f"Saved {template_path} — removed {len(removed)} shape(s).")
    return len(removed)


def main() -> int:
    """Entrypoint."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "template",
        type=Path,
        nargs="?",
        default=DEFAULT_TEMPLATE,
        help="Path to the .pptx template (defaults to the v1.6 GreenScout template).",
    )
    args = parser.parse_args()

    if not args.template.exists():
        print(f"ERROR: template not found at {args.template}", file=sys.stderr)
        return 2

    result = remove_slide4_image(args.template)
    if result < 0:
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main())
